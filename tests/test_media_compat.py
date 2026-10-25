"""Exercise the media paths affected by the MoviePy/Pillow version override."""

from base64 import b64decode
from io import BytesIO

import numpy as np
import pytest
from moviepy import ColorClip, CompositeVideoClip, ImageClip, VideoFileClip
from moviepy.video.fx import Resize, Rotate
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from stirrup.core.models import ImageContentBlock, VideoContentBlock


@pytest.mark.parametrize("mode", ["RGB", "RGBA", "L", "P"])
def test_stirrup_image_validation_and_downscaling(mode):
    source = BytesIO()
    Image.new(mode, (80, 40)).save(source, format="PNG")
    block = ImageContentBlock(data=source.getvalue())

    prefix, encoded = block.to_base64_url(max_pixels=800).split(",", 1)

    assert prefix == "data:image/png;base64"
    with Image.open(BytesIO(b64decode(encoded))) as result:
        result.load()
        assert result.mode == "RGB"
        assert result.size == (40, 20)


def test_moviepy_resize_rotate_and_alpha_composite():
    pixels = np.zeros((16, 24, 4), dtype=np.uint8)
    pixels[:, :, 0] = 255
    pixels[:, :, 3] = 128
    with ImageClip(pixels, duration=1) as source:
        resized = source.with_effects([Resize(new_size=(12, 8))])
        assert resized.get_frame(0).shape == (8, 12, 3)
        np.testing.assert_allclose(resized.mask.get_frame(0), 128 / 255, atol=0.01)
        rotated = resized.with_effects([Rotate(30)])
        assert rotated.get_frame(0).shape[2] == 3
        with CompositeVideoClip(
            [resized], size=(12, 8), bg_color=(0, 0, 255)
        ) as composite:
            frame = composite.get_frame(0)
            np.testing.assert_allclose(frame[4, 6], [128, 0, 127], atol=1)


def test_stirrup_video_transcode_and_downscale(tmp_path):
    source = tmp_path / "source.mp4"
    with ColorClip(size=(64, 32), color=(255, 0, 0), duration=0.5) as clip:
        clip.write_videofile(
            str(source), fps=4, codec="libx264", audio=False, logger=None
        )
    block = VideoContentBlock(data=source.read_bytes())

    prefix, encoded = block.to_base64_url(max_pixels=512, fps=4).split(",", 1)

    assert prefix == "data:video/mp4;base64"
    result = tmp_path / "resized.mp4"
    result.write_bytes(b64decode(encoded))
    with VideoFileClip(str(result)) as clip:
        assert clip.size == [32, 16]
        assert clip.fps == 4
        frame = clip.get_frame(0)
        assert frame.shape == (16, 32, 3)
        assert frame[:, :, 0].mean() > 240
        assert frame[:, :, 1:].mean() < 10


def test_pptx_image_embedding_round_trip():
    source = BytesIO()
    Image.new("RGB", (40, 20), color="red").save(source, format="PNG")
    source.seek(0)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(source, Inches(1), Inches(1))
    output = BytesIO()
    presentation.save(output)
    output.seek(0)

    result = Presentation(output).slides[0].shapes[0].image
    assert result.size == (40, 20)
    assert result.content_type == "image/png"
